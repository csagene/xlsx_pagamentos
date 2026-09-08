# -*- coding: utf-8 -*-
with open("app_dashboard.py", "r", encoding="utf-8") as f:
    content = f.read()

func_code = """
        def formatar_excel(writer, df_to_write, sheet_name, filtros):
            df_to_write.to_excel(writer, index=False, sheet_name=sheet_name)
            workbook = writer.book
            worksheet = writer.sheets[sheet_name]
            
            # Fixar cabeçalho
            worksheet.freeze_panes = 'A2'
            
            # Destacar linha de totais (a última linha)
            try:
                from openpyxl.styles import PatternFill, Font
                fill_totais = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")
                font_totais = Font(bold=True)
                
                max_row = worksheet.max_row
                for col in range(1, worksheet.max_column + 1):
                    cell = worksheet.cell(row=max_row, column=col)
                    cell.fill = fill_totais
                    cell.font = font_totais
            except:
                pass
                
            # Adicionar filtros numa aba separada se não existir
            if filtros and "Filtros" not in writer.sheets:
                import pandas as pd
                df_filtros = pd.DataFrame({"Filtros Aplicados": filtros})
                df_filtros.to_excel(writer, index=False, sheet_name="Filtros")
                try:
                    ws_filtros = writer.sheets["Filtros"]
                    ws_filtros.column_dimensions['A'].width = 100
                except:
                    pass

        tab1, tab2 = st.tabs(["?? Tabela 1 — PAGAMENTOS MENSAL", "?? Tabela 2 — PAGAMENTOS CUMULATIVO"])
"""

# Find the exact string to replace
old_tab_def = '        tab1, tab2 = st.tabs(["?? Tabela 1 — PAGAMENTOS MENSAL", "?? Tabela 2 — PAGAMENTOS CUMULATIVO"])'

content = content.replace(old_tab_def, func_code)

old_write1 = "                rel_display_completo.to_excel(writer, index=False, sheet_name='PAGAMENTOS_MENSAL')"
new_write1 = "                formatar_excel(writer, rel_display_completo, 'PAGAMENTOS_MENSAL', st.session_state.get('filtros_aplicados_texto', []))"
content = content.replace(old_write1, new_write1)

old_write2 = "                rel_cumul_display_completo.to_excel(writer, index=False, sheet_name='PAGAMENTOS_CUMULATIVO')"
new_write2 = "                formatar_excel(writer, rel_cumul_display_completo, 'PAGAMENTOS_CUMULATIVO', st.session_state.get('filtros_aplicados_texto', []))"
content = content.replace(old_write2, new_write2)

with open("app_dashboard.py", "w", encoding="utf-8") as f:
    f.write(content)

print("Success")
